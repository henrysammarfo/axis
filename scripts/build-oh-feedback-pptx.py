#!/usr/bin/env python3
"""AXIS Open House feedback deck — UXmaxx-grade craft (Pillow → pptx).

Story spine (compressed to 3 slides / ~90s):
  1 Hook + problem (human, brand-first)
  2 Product (English → basket) + honesty seal
  3 Depth + live URL + ask

Output:
  docs/pitch/slides/01-continuity.png
  docs/pitch/slides/02-baskets.png
  docs/pitch/slides/03-depth.png
  docs/pitch/AXIS-OH-Feedback-v2.pptx
"""

from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Emu

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "pitch"
SLIDES = OUT / "slides"
LOGOS = ROOT / "public" / "logos"
PUBLIC = ROOT / "public"

W, H = 1920, 1080
BLACK = (0, 0, 0)
WHITE = (255, 255, 255)
MUTED = (168, 168, 168)  # higher contrast than 140
SOFT = (110, 110, 110)
LIME = (200, 245, 74)
LINE = (48, 48, 48)
CARD = (14, 14, 14)

NOTES = [
    "I’m Henry. DeFi yield is real — normal people still bounce on MetaMask, gas, and sign-every-trade. AXIS already won UXmaxx on Arbitrum: Google login, Set.Forget.Earn, zero signing after login. That’s live money today.",
    "For Open House we put US stock tokens under that same Google door. You say ‘tech yes, oil no’ — AXIS builds a Robinhood Chain practice basket. Testnet only. Plan is not a fill. Live Arbitrum yield stays on a separate tab — no loophole.",
    "Judges score attract and retain, so we shipped depth behind a simple path: instrument truth, labeled USDG rail, EU/APAC beachhead, packages. Walk it at axis-mainnet.vercel.app/demo/baskets. What still feels thin? I’m here for notes.",
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        "C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def load_rgba(path: Path, size: int) -> Image.Image | None:
    if not path.exists():
        return None
    im = Image.open(path).convert("RGBA")
    im.thumbnail((size, size), Image.Resampling.LANCZOS)
    canvas = Image.new("RGBA", (size, size), (0, 0, 0, 0))
    canvas.paste(im, ((size - im.width) // 2, (size - im.height) // 2), im)
    return canvas


def wash(im: Image.Image) -> None:
    draw = ImageDraw.Draw(im)
    for i in range(260):
        a = int(32 * (1 - i / 260))
        if a <= 0:
            continue
        c = (min(255, LIME[0] * a // 255), min(255, LIME[1] * a // 255), min(255, LIME[2] * a // 255))
        draw.line([(0, i), (W, i)], fill=c)


def footer(draw: ImageDraw.ImageDraw, n: int, total: int = 3) -> None:
    draw.line([(80, H - 78), (W - 80, H - 78)], fill=LINE, width=1)
    draw.text((80, H - 56), "AXIS", font=font(22, True), fill=MUTED)
    draw.text((W - 130, H - 56), f"{n} / {total}", font=font(22), fill=MUTED)


def partner_row(im: Image.Image, draw: ImageDraw.ImageDraw, y: int) -> None:
    items = [
        (LOGOS / "chains" / "arbitrum.png", "ARBITRUM"),
        (LOGOS / "chains" / "robinhood.png", "ROBINHOOD"),
    ]
    x = 80
    for path, label in items:
        logo = load_rgba(path, 48)
        if logo:
            # soft chip
            chip = Image.new("RGBA", (56, 56), (255, 255, 255, 18))
            im.paste(Image.alpha_composite(Image.new("RGBA", (56, 56), (0, 0, 0, 0)), chip), (x, y))
            im.paste(logo, (x + 4, y + 4), logo)
        draw.text((x + 68, y + 14), label, font=font(18), fill=MUTED)
        x += 280


def slide_1() -> Image.Image:
    im = Image.new("RGB", (W, H), BLACK)
    wash(im)
    draw = ImageDraw.Draw(im)

    draw.text((80, 64), "OPEN HOUSE  ·  90-SECOND FEEDBACK", font=font(20), fill=LIME)
    draw.text((80, 160), "AXIS", font=font(128, True), fill=WHITE)

    draw.text((80, 340), "Yield is real.", font=font(56, True), fill=WHITE)
    draw.text((80, 420), "Signing is the wall.", font=font(56, True), fill=MUTED)

    draw.text(
        (80, 540),
        "MetaMask · gas · bridges · sign every trade.\n"
        "AXIS already ships the other door on Arbitrum:\n"
        "Google in → Set.Forget.Earn → zero signing after login.",
        font=font(28),
        fill=SOFT,
    )

    partner_row(im, draw, 760)
    draw.text((80, 860), "UXmaxx Arbitrum bounty winner  ·  Live on mainnet", font=font(22), fill=LIME)

    footer(draw, 1)
    return im


def slide_2() -> Image.Image:
    im = Image.new("RGB", (W, H), BLACK)
    wash(im)
    draw = ImageDraw.Draw(im)

    draw.text((80, 64), "PRODUCT  ·  SAME GOOGLE DOOR", font=font(20), fill=LIME)
    draw.text((80, 120), "English → practice stock basket", font=font(48, True), fill=WHITE)
    draw.text(
        (80, 190),
        "“tech yes, oil no”  ·  fail-closed  ·  plan ≠ fill  ·  testnet sealed from mainnet",
        font=font(24),
        fill=MUTED,
    )

    # Prompt card
    draw.rounded_rectangle([80, 270, 1100, 390], radius=20, fill=CARD, outline=LINE, width=2)
    draw.text((110, 292), "YOUR VIBE", font=font(16), fill=SOFT)
    draw.text((110, 330), "tech yes, oil no", font=font(34, True), fill=WHITE)

    legs = [("TSLA", "22%"), ("NVDA", "22%"), ("AAPL", "18%"), ("MSFT", "18%"), ("AMZN", "20%")]
    x = 80
    for sym, wgt in legs:
        draw.rounded_rectangle([x, 420, x + 190, 600], radius=16, fill=CARD, outline=LINE, width=2)
        logo = load_rgba(LOGOS / "stocks" / f"{sym.lower()}.png", 40)
        if logo:
            im.paste(logo, (x + 24, 450), logo)
        draw.text((x + 78, 458), sym, font=font(20, True), fill=WHITE)
        draw.text((x + 24, 530), wgt, font=font(30, True), fill=LIME)
        x += 208

    # Seal strip
    draw.rounded_rectangle([80, 640, 1100, 760], radius=16, fill=(28, 24, 8), outline=(120, 90, 30), width=2)
    draw.text((110, 668), "NETWORK SEAL", font=font(16), fill=LIME)
    draw.text(
        (110, 702),
        "Live money = Arbitrum Overview   ·   Practice stocks = RH testnet 46630 only",
        font=font(24),
        fill=WHITE,
    )

    # Right honesty
    draw.text((1200, 280), "Honesty", font=font(24, True), fill=LIME)
    bullets = [
        "Public testnet — no cash value",
        "Oil blocked by BasketPolicy",
        "Never invent fills / silent OFT",
        "Complexity under Advanced",
        "Newcomers see 3 steps only",
    ]
    y = 340
    for b in bullets:
        draw.ellipse([1200, y + 10, 1214, y + 24], fill=LIME)
        draw.text((1234, y), b, font=font(24), fill=WHITE)
        y += 58

    footer(draw, 2)
    return im


def slide_3() -> Image.Image:
    im = Image.new("RGB", (W, H), BLACK)
    wash(im)
    draw = ImageDraw.Draw(im)

    draw.text((80, 64), "DEPTH  ·  ASK", font=font(20), fill=LIME)
    draw.text((80, 120), "Simple on the surface.", font=font(48, True), fill=WHITE)
    draw.text((80, 190), "Truth underneath.", font=font(48, True), fill=MUTED)

    cells = [
        ("01", "Instrument truth", "Same ticker ≠ same claim"),
        ("02", "USDG rail", "Labeled · AXIS does not bridge"),
        ("03", "Beachhead", "EU / APAC waitlist"),
        ("04", "Packages", "Catalog only — no fake checkout"),
    ]
    positions = [(80, 290), (980, 290), (80, 500), (980, 500)]
    for (n, title, sub), (x, y) in zip(cells, positions):
        draw.rounded_rectangle([x, y, x + 860, y + 170], radius=20, fill=CARD, outline=LINE, width=2)
        draw.text((x + 36, y + 28), n, font=font(22), fill=LIME)
        draw.text((x + 36, y + 68), title, font=font(32, True), fill=WHITE)
        draw.text((x + 36, y + 118), sub, font=font(22), fill=MUTED)

    draw.text((80, 720), "Public depth (no login)", font=font(20), fill=SOFT)
    draw.text(
        (80, 760),
        "axis-mainnet.vercel.app/demo/baskets",
        font=font(32, True),
        fill=LIME,
    )
    draw.text(
        (80, 830),
        "What still feels thin? Tell me — I’m here for notes.",
        font=font(26),
        fill=MUTED,
    )

    # tiny brand mark bottom-right area before footer
    fav = PUBLIC / "favicon.png"
    if not fav.exists():
        fav = PUBLIC / "apple-touch-icon.png"
    mark = load_rgba(fav, 40)
    if mark:
        im.paste(mark, (W - 160, 760), mark)

    footer(draw, 3)
    return im


def pack_pptx(paths: list[Path]) -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    for i, path in enumerate(paths):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(path), Emu(0), Emu(0), width=prs.slide_width, height=prs.slide_height)
        slide.notes_slide.notes_text_frame.text = NOTES[i]

    out = OUT / "AXIS-OH-Feedback-v2.pptx"
    try:
        prs.save(out)
    except PermissionError:
        out = OUT / "AXIS-OH-Feedback-v3.pptx"
        prs.save(out)
    return out


def main() -> None:
    SLIDES.mkdir(parents=True, exist_ok=True)
    builders = [slide_1, slide_2, slide_3]
    names = ["01-continuity.png", "02-baskets.png", "03-depth.png"]
    paths: list[Path] = []
    for build, name in zip(builders, names):
        im = build()
        path = SLIDES / name
        im.save(path, "PNG", optimize=True)
        paths.append(path)
        print("wrote", path)
    pptx = pack_pptx(paths)
    print("wrote", pptx)


if __name__ == "__main__":
    main()
